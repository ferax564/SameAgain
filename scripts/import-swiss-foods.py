"""Reproducible FSVO generic-food importer. Usage: python scripts/import-swiss-foods.py workbook.xlsx
Download source: https://naehrwertdaten.ch/wp-content/uploads/2026/07/Swiss_food_composition_database.xlsx
Requires openpyxl. Specific dataset reuse permission: https://valeursnutritives.ch/en/downloads/
Missing, below-detection and nonnumeric values remain unknown, not zero.
"""
import sys,json,hashlib,datetime
from pathlib import Path
from openpyxl import load_workbook
p=Path(sys.argv[1]); w=load_workbook(p,read_only=True,data_only=True)
cols={'energy-kcal':11,'fat':14,'saturated-fat':17,'carbohydrates':41,'sugars':44,'fiber':50,'proteins':53,'salt':56,'vitamin-a':65,'vitamin-b1':80,'vitamin-b2':83,'vitamin-b6':86,'vitamin-b12':89,'vitamin-b9':98,'vitamin-c':104,'vitamin-d':107,'vitamin-e':110,'potassium':113,'sodium':116,'calcium':122,'magnesium':125,'phosphorus':128,'iron':131,'iodine':134,'zinc':137,'selenium':140}
foods=[]; provenance={}
for r in w['Generic Foods'].iter_rows(min_row=4,values_only=True):
 if not isinstance(r[0],(int,float)) or not r[3]:continue
 if r[7]!='per 100g edible portion':continue
 n={k:r[i] for k,i in cols.items() if isinstance(r[i],(int,float)) and r[i]>=0}
 fid='swiss:'+str(int(r[0])); foods.append({'id':fid,'name':r[3],'categories':str(r[5] or '').split(';'),'countries':[],'nutrition':n,'basis':'100g','source':'Swiss Food Composition Database · FSVO · v7.1 (generic food)','sourceUrl':'https://valeursnutritives.ch/en/downloads/','retrieved':1788566400000})
 provenance[fid]={'synonyms':r[4],'changed':str(r[143] or ''),'nutrients':{k:{'value':r[i],'derivation':r[i+1],'source':r[i+2]} for k,i in cols.items()}}
root=Path(__file__).resolve().parents[1]
(root/'lib/swiss-foods.json').write_text(json.dumps(foods,ensure_ascii=False,separators=(',',':')))
(root/'lib/swiss-provenance.json').write_text(json.dumps({'version':'7.1 (01.07.2026)','sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'attribution':'Federal Food Safety and Veterinary Office FSVO, Swiss Food Composition Database','permission':'https://valeursnutritives.ch/en/downloads/','sources':{str(int(r[0])):r[1] for r in w['Sources'].iter_rows(min_row=4,values_only=True) if isinstance(r[0],(int,float))},'records':provenance},ensure_ascii=False,default=str,separators=(',',':')))
print(f'Imported {len(foods)} generic foods, with explicit mass basis and nutrient provenance.')
